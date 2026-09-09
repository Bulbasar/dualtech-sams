from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
import re

def clean_text(text):
    if not text:
        return ""
    # Remove extra spaces
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def create_pptx(html_path, pptx_path):
    with open(html_path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')

    prs = Presentation()

    # Layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]

    slides_wrapper = soup.find('div', id='slidesWrapper')
    if not slides_wrapper:
        print("No slides wrapper found.")
        return

    slides = slides_wrapper.find_all('div', class_=re.compile(r'\bslide\b'), recursive=False)

    for i, slide_div in enumerate(slides):
        title_el = slide_div.find(['h1', 'h2'], class_='slide-title')
        subtitle_el = slide_div.find('p', class_='slide-subtitle')
        
        # Decide layout
        if i == 0:
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = clean_text(title_el.get_text()) if title_el else "Slide Title"
            subtitle.text = clean_text(subtitle_el.get_text()) if subtitle_el else ""
        else:
            slide = prs.slides.add_slide(bullet_slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]

            title_shape.text = clean_text(title_el.get_text()) if title_el else "Slide Title"
            
            tf = body_shape.text_frame
            subtitle_text = clean_text(subtitle_el.get_text()) if subtitle_el else ""
            if subtitle_text:
                tf.text = subtitle_text

            # Extract features/bullets if any
            feat_cards = slide_div.find_all('div', class_='feat-card')
            for card in feat_cards:
                feat_name = card.find('div', class_='feat-name')
                feat_desc = card.find('p', class_='feat-desc')
                
                if feat_name:
                    p = tf.add_paragraph() if tf.text else tf
                    p.text = clean_text(feat_name.get_text())
                    p.level = 0 if not tf.text else 1
                    
                    if feat_desc:
                        p2 = tf.add_paragraph()
                        p2.text = clean_text(feat_desc.get_text())
                        p2.level = 1 if not tf.text else 2

            check_lists = slide_div.find_all('ul', class_='check-list')
            for ul in check_lists:
                for li in ul.find_all('li'):
                    text_parts = []
                    for child in li.children:
                        if hasattr(child, 'get_text'):
                            text = child.get_text()
                        elif isinstance(child, str):
                            text = child
                        else:
                            text = ""
                        # filter out icons
                        if "check-icon" not in getattr(child, 'get', lambda x: [])('class', []):
                            text_parts.append(clean_text(text))
                            
                    li_text = " ".join([t for t in text_parts if t]).replace("  ", " ").strip()
                    if li_text:
                        p = tf.add_paragraph() if tf.text else tf
                        p.text = li_text
                        p.level = 0 if not tf.text else 1

            kpi_rows = slide_div.find_all('div', class_='kpi-row')
            for row in kpi_rows:
                for kpi in row.find_all('div', class_='kpi'):
                    val = kpi.find('div', class_='kpi-value')
                    lbl = kpi.find('div', class_='kpi-label')
                    if val and lbl:
                        p = tf.add_paragraph() if tf.text else tf
                        p.text = f"{clean_text(lbl.get_text())}: {clean_text(val.get_text())}"
                        p.level = 0 if not tf.text else 1

            # Handle step-cards if any
            step_cards = slide_div.find_all('div', class_='step-card')
            for card in step_cards:
                title = card.find('div', class_='step-title')
                desc = card.find('p', class_='step-desc')
                if title:
                    p = tf.add_paragraph() if tf.text else tf
                    p.text = clean_text(title.get_text())
                    p.level = 0 if not tf.text else 1
                    if desc:
                        p2 = tf.add_paragraph()
                        p2.text = clean_text(desc.get_text())
                        p2.level = 1 if not tf.text else 2

    prs.save(pptx_path)
    print(f"Presentation saved to {pptx_path}")

if __name__ == '__main__':
    create_pptx('presentation.html', 'presentation.pptx')
